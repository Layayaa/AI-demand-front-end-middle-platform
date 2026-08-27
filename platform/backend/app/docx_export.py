import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def render_document_docx(document: dict[str, Any], project: dict[str, Any]) -> BytesIO:
    output = BytesIO()
    word = Document()
    styles = word.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)

    title = word.add_heading(str(document.get("title") or project.get("title") or "需求方案"), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = word.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"项目：{project.get('title') or '-'}  |  版本：V{document.get('version') or 1}"
    ).italic = True

    for raw_line in str(document.get("markdown_content") or "").splitlines():
        line = raw_line.strip()
        if not line or line == "---":
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            word.add_heading(_plain_text(heading.group(2)), level=min(len(heading.group(1)), 3))
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            _add_markdown_runs(word.add_paragraph(style="List Bullet"), bullet.group(1))
            continue
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if numbered:
            _add_markdown_runs(word.add_paragraph(style="List Number"), numbered.group(1))
            continue
        _add_markdown_runs(word.add_paragraph(), line)

    section = word.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AI 需求前置分析中台 · 生成文档")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    word.save(output)
    output.seek(0)
    return output


def safe_docx_filename(value: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in "-_. " else "_" for char in value).strip()
    return f"{cleaned or '需求方案'}.docx"


def _plain_text(value: str) -> str:
    return re.sub(r"[`*_]", "", value).strip()


def _add_markdown_runs(paragraph: Any, value: str) -> None:
    cursor = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|`(.+?)`", value):
        if match.start() > cursor:
            paragraph.add_run(value[cursor:match.start()])
        run = paragraph.add_run(match.group(1) or match.group(2) or "")
        run.bold = bool(match.group(1))
        if match.group(2):
            run.font.name = "Menlo"
        cursor = match.end()
    if cursor < len(value):
        paragraph.add_run(value[cursor:])
