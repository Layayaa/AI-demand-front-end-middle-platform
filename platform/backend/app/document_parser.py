from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".docx",
    ".pdf",
    ".txt",
    ".md",
    ".xlsx",
    ".xls",
    ".csv",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
    ".bmp",
}


class DocumentParseError(RuntimeError):
    """文件可读取但无法按当前支持格式解析。"""


def parser_type_for(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentParseError(f"暂不支持该文件类型，仅支持：{allowed}")
    return suffix.lstrip(".")


def parse_document(path: Path, original_filename: str) -> dict:
    parser_type = parser_type_for(original_filename)
    if parser_type == "docx":
        return _parse_docx(path)
    if parser_type == "pdf":
        return _parse_pdf(path)
    if parser_type in {"txt", "md"}:
        return _parse_text(path, parser_type)
    if parser_type == "xlsx":
        return _parse_xlsx(path)
    if parser_type == "xls":
        return _parse_xls(path)
    if parser_type == "csv":
        return _parse_text(path, parser_type)
    if parser_type in {"png", "jpg", "jpeg", "webp", "gif", "bmp"}:
        return _parse_image(path, parser_type)
    raise DocumentParseError("暂不支持该文件类型")


def _parse_docx(path: Path) -> dict:
    try:
        document = Document(path)
    except Exception as exc:
        raise DocumentParseError(f"DOCX 解析失败：{exc}") from exc

    paragraphs = [
        f"[段落 {index}]\n{paragraph.text.strip()}"
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    table_blocks: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            table_blocks.append(f"[表格 {table_index}]\n" + "\n".join(rows))

    return {
        "parser_type": "docx",
        "page_count": None,
        "text": "\n\n".join([*paragraphs, *table_blocks]).strip(),
        "metadata": {
            "paragraphCount": len(paragraphs),
            "tableCount": len(table_blocks),
        },
    }


def _parse_pdf(path: Path) -> dict:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentParseError(f"PDF 解析失败：{exc}") from exc

    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[第 {index} 页]\n{text}")

    return {
        "parser_type": "pdf",
        "page_count": len(reader.pages),
        "text": "\n\n".join(pages).strip(),
        "metadata": {
            "pageCount": len(reader.pages),
            "textPageCount": len(pages),
        },
    }


def _parse_text(path: Path, parser_type: str) -> dict:
    try:
        text = path.read_text(encoding="utf-8")
        encoding = "utf-8"
    except UnicodeDecodeError:
        text = path.read_text(encoding="gb18030", errors="replace")
        encoding = "gb18030"

    return {
        "parser_type": parser_type,
        "page_count": None,
        "text": text.strip(),
        "metadata": {
            "encoding": encoding,
            "lineCount": len(text.splitlines()),
        },
    }


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace("\n", " ")


def _rows_to_text(sheet_name: str, rows: list[list[object]]) -> str:
    lines = []
    for row in rows:
        cells = [_cell_text(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            lines.append(" | ".join(cells))
    if not lines:
        return ""
    return f"[工作表：{sheet_name}]\n" + "\n".join(lines)


def _parse_xlsx(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("XLSX 解析依赖未安装，请安装 openpyxl") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        sheet_metadata = []
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            block = _rows_to_text(worksheet.title, rows)
            if block:
                blocks.append(block)
            sheet_metadata.append({"name": worksheet.title, "rowCount": len(rows)})
        workbook.close()
    except Exception as exc:
        raise DocumentParseError(f"XLSX 解析失败：{exc}") from exc

    return {
        "parser_type": "xlsx",
        "page_count": None,
        "text": "\n\n".join(blocks).strip(),
        "metadata": {
            "sheetCount": len(sheet_metadata),
            "sheets": sheet_metadata,
        },
    }


def _parse_xls(path: Path) -> dict:
    try:
        import xlrd
    except ImportError as exc:
        raise DocumentParseError("XLS 解析依赖未安装，请安装 xlrd") from exc

    try:
        workbook = xlrd.open_workbook(path, on_demand=True)
        blocks = []
        sheet_metadata = []
        for sheet_name in workbook.sheet_names():
            worksheet = workbook.sheet_by_name(sheet_name)
            rows = [
                [worksheet.cell_value(row_index, column_index) for column_index in range(worksheet.ncols)]
                for row_index in range(worksheet.nrows)
            ]
            block = _rows_to_text(sheet_name, rows)
            if block:
                blocks.append(block)
            sheet_metadata.append({"name": sheet_name, "rowCount": worksheet.nrows})
        workbook.release_resources()
    except Exception as exc:
        raise DocumentParseError(f"XLS 解析失败：{exc}") from exc

    return {
        "parser_type": "xls",
        "page_count": None,
        "text": "\n\n".join(blocks).strip(),
        "metadata": {
            "sheetCount": len(sheet_metadata),
            "sheets": sheet_metadata,
        },
    }


def _parse_image(path: Path, parser_type: str) -> dict:
    try:
        file_size = path.stat().st_size
    except OSError:
        file_size = 0
    return {
        "parser_type": parser_type,
        "page_count": None,
        "text": "",
        "metadata": {
            "mediaType": "image",
            "fileSize": file_size,
            "textExtraction": "pending_vision_ocr",
            "message": "图片已保存；当前版本暂未启用图片文字识别。",
        },
    }
