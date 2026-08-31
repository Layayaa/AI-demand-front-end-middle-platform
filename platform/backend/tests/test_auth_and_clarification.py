import base64
import httpx
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from app.auth import (
    AuthenticationError,
    create_access_token,
    decode_access_token,
    hash_password,
    verify_password,
)
from app.clarification import (
    apply_ai_direction,
    apply_answer,
    confirmation_gaps_for,
    extract_minutes,
    extract_people,
    extract_value,
    gaps_for,
    new_profile,
    next_message,
    questions_for,
    visible_gaps_for,
    _next_unanswered_question,
    split_steps,
)
from app.config import Settings
from app.document_parser import parse_document, parser_type_for
from app.docx_export import render_document_docx
from app.knowledge import collect_plan_references
from app.material_profile import (
    add_material_candidates,
    merge_material_into_profile,
    resolve_material_candidate,
)
from app.main import _requester_document
from app.assessment import assess, merge_model_assessment
from app.llm import (
    _parse_requirement_assessment,
    _partial_json_string_field,
    _project_materials_text,
    _single_question,
)
from app.plans import generate_documents, generate_requirement_explanation


class AuthTests(unittest.TestCase):
    def test_password_hash_and_signed_token(self) -> None:
        password_hash = hash_password("local-password")
        self.assertTrue(verify_password("local-password", password_hash))
        self.assertFalse(verify_password("incorrect-password", password_hash))

        token = create_access_token(
            {"id": "user-1", "role": "requester"},
            "test-secret",
            30,
        )
        payload = decode_access_token(token, "test-secret")
        self.assertEqual(payload["sub"], "user-1")
        self.assertEqual(payload["role"], "requester")
        with self.assertRaises(AuthenticationError):
            decode_access_token(token, "different-secret")


class ClarificationTests(unittest.TestCase):
    def test_continue_request_from_confirmation_moves_to_a_real_gap(self) -> None:
        profile = new_profile(
            title="退款审核", department="客服", summary="", requirement_type="sop"
        )
        profile["stage"] = "confirm"
        profile["aiJudged"] = True
        profile["businessContext"]["businessObject"] = "退款申请"
        result = apply_answer(
            profile,
            "这部分先按当前口径执行；如果会影响方案设计，请继续指出需要明确的边界。",
        )
        self.assertEqual(result["reason"], "continue_requested")
        self.assertEqual(result["profile"]["stage"], "clarifying")
        self.assertNotEqual(result["profile"]["activeQuestion"]["key"], "other")

    def test_bare_number_is_accepted_for_people(self) -> None:
        self.assertEqual(extract_people("3"), 3)
        self.assertEqual(extract_people("3人"), 3)

    def test_multi_sentence_step_answer_is_split_into_steps(self) -> None:
        answer = (
            "系统同步订单和物流数据。"
            "系统校验退款金额与订单状态。"
            "高金额或证据不足的工单转人工复核。"
        )
        self.assertEqual(len(split_steps(answer)), 3)
        profile = new_profile(
            title="退款审核", department="客服", summary="", requirement_type="sop"
        )
        profile["activeQuestion"] = {"key": "steps", "content": "请描述流程", "chips": []}
        result = apply_answer(profile, answer)
        self.assertEqual(len(result["profile"]["sop"]["steps"]), 3)

    def _answer_business_context(self, profile, answers):
        for answer in answers:
            profile = apply_answer(profile, answer)["profile"]
        return profile

    def test_ai_led_profile_does_not_follow_fixed_question_order(self) -> None:
        profile = new_profile(
            title="滞销 SKU 处置判断",
            department="商品运营",
            summary="决定滞销库存该继续销售还是清仓。",
            requirement_type="decision",
        )
        first_prompt = next_message(profile)
        self.assertEqual(first_prompt["stage"], "clarifying")
        self.assertEqual(first_prompt["questionKey"], "other")
        self.assertEqual(profile["businessContext"]["businessChain"], "")

        profile = apply_ai_direction(
            profile,
            {
                "status": "continue",
                "requirementType": "decision",
                "questionKey": "purpose",
                "question": "如果继续销售或清仓判断失误，最先会造成什么业务后果？",
                "facts": {
                    "businessChain": "商品、库存与商品后台运营",
                    "businessObject": "滞销 SKU 的库存处置",
                    "scope": "先覆盖直营网店常规商品，不含预售和跨境商品",
                },
                "criticalGaps": [],
            },
        )
        self.assertEqual(next_message(profile)["questionKey"], "purpose")
        self.assertEqual(profile["businessContext"]["businessChain"], "inventory_operations")
        self.assertEqual(profile["businessContext"]["businessObject"], "滞销 SKU 的库存处置")

        profile = apply_answer(profile, "减少库存资金占用，避免积压和仓储成本继续增加")["profile"]
        profile = apply_ai_direction(
            profile,
            {
                "status": "continue",
                "requirementType": "decision",
                "questionKey": "approvalBoundary",
                "question": "哪些处置结果必须交给主管确认？",
                "criticalGaps": [],
            },
        )
        self.assertEqual(next_message(profile)["questionKey"], "approvalBoundary")

    def test_ai_can_choose_business_context_question_without_starting_a_checklist(self) -> None:
        profile = new_profile(
            title="商品 VOC 优化建议",
            department="商品运营",
            summary="整理客服反馈和销售数据，输出优化建议。",
            requirement_type="sop",
        )

        first_prompt = next_message(profile)
        self.assertEqual(first_prompt["questionKey"], "other")

        profile = apply_ai_direction(
            profile,
            {
                "status": "continue",
                "requirementType": "sop",
                "questionKey": "businessObject",
                "question": "你最想先解决哪类客户反馈，以及希望它最终改变什么？",
                "facts": {"businessChain": "商品优化与客户 VOC"},
                "criticalGaps": [],
            },
        )

        self.assertEqual(profile["businessContext"]["businessChain"], "product_voc")
        self.assertIn("客服反馈与销售分析场景", profile["businessContext"]["candidateSignals"])
        self.assertEqual(next_message(profile)["questionKey"], "businessObject")

    def test_ai_can_prioritize_a_question_outside_the_fixed_sop_order(self) -> None:
        profile = new_profile(
            title="售后退款自动审核",
            department="客服中心",
            summary="希望根据订单、物流和退款原因自动审核退款。",
            requirement_type="sop",
        )
        profile = apply_ai_direction(
            profile,
            {
                "status": "continue",
                "requirementType": "sop",
                "questionKey": "approvalBoundary",
                "question": "哪些退款必须保留给人工复核，不能由系统直接处理？",
                "chips": ["高金额退款", "证据不足", "疑似异常退款"],
                "criticalGaps": [
                    {
                        "key": "approvalBoundary",
                        "label": "人工复核边界",
                        "reason": "直接决定自动审核的风险和权限设计。",
                    }
                ],
            },
        )

        prompt = next_message(profile)
        self.assertEqual(prompt["questionKey"], "approvalBoundary")
        self.assertEqual(prompt["content"], "哪些退款必须保留给人工复核，不能由系统直接处理？")

        profile = apply_answer(profile, "高金额、证据不足和疑似异常退款都交客服复核")["profile"]
        self.assertEqual(
            profile["businessContext"]["approvalBoundary"],
            "高金额、证据不足和疑似异常退款都交客服复核",
        )

    def test_rule_fallback_prioritizes_automation_boundary_over_efficiency_fields(self) -> None:
        profile = new_profile(
            title="库存自动同步", department="商品运营", summary="每天同步库存", requirement_type="sop"
        )
        profile["businessContext"].update(
            {
                "businessChain": "inventory_operations",
                "businessObject": "SKU 库存",
                "scope": "直营网店常规商品",
            }
        )
        question = _next_unanswered_question(profile, excluded=set())
        self.assertEqual(question["key"], "approvalBoundary")

    def test_time_range_answer_is_recorded_and_does_not_repeat_question(self) -> None:
        profile = new_profile(
            title="自动跑库存流程",
            department="商品运营",
            summary="希望减少人工处理库存流程的时间。",
            requirement_type="sop",
        )
        profile["businessContext"].update(
            {
                "businessChain": "inventory_operations",
                "businessObject": "库存设置",
                "scope": "先覆盖直营网店常规商品",
                "currentProcess": "运营每天在商品后台和 Excel 之间人工处理",
                "approvalBoundary": "数据异常时转人工",
            }
        )
        profile["sop"].update(
            {
                "owner": "运营主管",
                "purpose": "缩短库存处理时间",
            }
        )
        profile["activeQuestion"] = {
            "key": "durationMin",
            "content": "现在跑完整条流程大约需要多少分钟？",
            "chips": [],
        }

        result = apply_answer(profile, "4-6个小时")
        profile = result["profile"]

        self.assertTrue(result["changed"])
        self.assertEqual(profile["sop"]["durationMin"], 300)
        self.assertEqual(extract_minutes("45分钟"), 45)
        self.assertEqual(extract_minutes("4小时"), 240)
        self.assertEqual(extract_minutes("半小时"), 30)
        self.assertEqual(extract_minutes("4-6小时"), 300)

        profile = apply_ai_direction(
            profile,
            {
                "status": "continue",
                "requirementType": "sop",
                "questionKey": "durationMin",
                "question": "现在跑完整条流程大约需要多少分钟？",
                "chips": [],
            },
        )
        self.assertNotEqual(next_message(profile)["questionKey"], "durationMin")

    def test_repeated_freeform_question_moves_to_an_unanswered_topic(self) -> None:
        profile = new_profile(
            title="库存预测",
            department="计划部",
            summary="根据库存和销售数据辅助补货。",
            requirement_type="sop",
        )
        first_direction = {
            "status": "continue",
            "requirementType": "sop",
            "questionKey": "other",
            "question": "当前最容易出错的地方是什么？",
            "chips": [],
        }
        profile = apply_ai_direction(profile, first_direction)
        profile = apply_ai_direction(
            profile,
            first_direction,
            recent_question_texts=["当前最容易出错的地方是什么？"],
        )
        self.assertNotEqual(
            next_message(profile)["content"],
            "当前最容易出错的地方是什么？",
        )

    def test_value_prompt_collects_impact_evidence_instead_of_forcing_a_level(self) -> None:
        profile = new_profile(
            title="库存预测",
            department="计划部",
            summary="根据库存和销售数据辅助补货。",
            requirement_type="sop",
        )
        value_question = next(
            item for item in questions_for("sop") if item["key"] == "value"
        )
        self.assertNotIn("高、中还是低", value_question["prompt"])
        self.assertTrue(value_question["ai_only"])

        profile["activeQuestion"] = {
            "key": "value",
            "content": value_question["prompt"],
            "chips": list(value_question["chips"]),
        }
        result = apply_answer(profile, "可能影响库存周转和断货风险，但暂时没有金额数据")
        value = result["profile"]["sop"]["value"]

        self.assertTrue(result["changed"])
        self.assertEqual(value["level"], "待模型判断")
        self.assertEqual(value["source"], "requester_evidence")
        self.assertIn("断货风险", value["reason"])
        self.assertEqual(extract_value("高")["source"], "requester_self_assessment")

    def test_fallback_value_score_does_not_treat_self_report_as_business_proof(self) -> None:
        profile = new_profile(
            title="一个内部报表",
            department="运营",
            summary="希望减少人工整理时间。",
            requirement_type="sop",
        )
        profile["sop"]["value"] = {
            "level": "高",
            "reason": "高",
            "source": "requester_self_assessment",
        }
        fallback = assess(profile, [])
        value_dimension = next(
            item for item in fallback["dimensions"] if item["name"] == "业务影响证据"
        )
        self.assertLess(value_dimension["score"], 70)

    def test_ai_can_stop_clarification_before_every_checklist_field_is_filled(self) -> None:
        profile = new_profile(
            title="售后退款自动审核",
            department="客服中心",
            summary="希望根据订单、物流和退款原因自动审核退款。",
            requirement_type="sop",
        )
        profile["businessContext"].update(
            {
                "businessChain": "order_finance",
                "businessObject": "退款申请自动审核",
                "scope": "先覆盖国内直营网店普通退款",
                "currentProcess": "客服在订单后台核对订单和物流后手工审核",
                "approvalBoundary": "高金额和疑似异常退款转人工",
            }
        )
        profile["sop"].update(
            {
                "owner": "客服主管",
                "purpose": "缩短退款审核时效",
                "steps": [{"seq": 1, "action": "核对订单和物流信息"}],
            }
        )
        profile = apply_ai_direction(
            profile,
            {
                "status": "ready_to_confirm",
                "requirementType": "sop",
                "questionKey": None,
                "question": "",
                "chips": [],
                "criticalGaps": [],
            },
        )

        self.assertEqual(profile["stage"], "confirm")
        self.assertEqual(next_message(profile)["stage"], "confirm")
        self.assertLess(profile["completeness"], 100)
        self.assertEqual(visible_gaps_for(profile), [])

    def test_ready_to_confirm_is_blocked_when_minimum_evidence_is_missing(self) -> None:
        profile = new_profile(
            title="唯品会",
            department="运营",
            summary="在唯品会网站上抓取价格标签。",
            requirement_type=None,
        )
        profile["businessContext"].update(
            {
                "businessChain": "inventory_operations",
                "businessObject": "全量公开商品",
                "scope": "全量公开商品",
            }
        )
        profile = apply_ai_direction(
            profile,
            {
                "status": "ready_to_confirm",
                "requirementType": "unknown",
                "questionKey": None,
                "question": "",
                "chips": [],
                "criticalGaps": [],
            },
        )

        self.assertEqual(profile["stage"], "clarifying")
        self.assertEqual(next_message(profile)["questionKey"], "other")
        self.assertIn("业务判断", next_message(profile)["content"])
        self.assertTrue(confirmation_gaps_for(profile))

    def test_business_object_answer_is_not_treated_as_business_chain(self) -> None:
        profile = new_profile(
            title="唯品会",
            department="运营",
            summary="在唯品会网站上抓取价格标签。",
            requirement_type=None,
        )
        result = apply_answer(profile, "全量公开商品")
        profile = result["profile"]

        self.assertTrue(result["changed"])
        self.assertEqual(result["reason"], "freeform")
        self.assertEqual(profile["businessContext"]["businessChain"], "")
        self.assertIn("全量公开商品", profile["notes"])

    def test_type_answer_during_business_chain_question_does_not_pollute_object(self) -> None:
        profile = new_profile(
            title="周预测表", department="计划部", summary="", requirement_type=None
        )
        profile["activeQuestion"] = {
            "key": "businessChain",
            "content": "请选择业务链路",
            "chips": [],
        }
        result = apply_answer(profile, "按步骤跑一条流程")
        self.assertEqual(result["reason"], "business_chain_unclear")
        self.assertEqual(result["profile"]["type"], "sop")
        self.assertEqual(result["profile"]["businessContext"]["businessObject"], "")
        self.assertEqual(result["profile"]["activeQuestion"]["key"], "businessChain")

    def test_confirm_button_cannot_bypass_missing_minimum_evidence(self) -> None:
        profile = new_profile(
            title="唯品会",
            department="运营",
            summary="在唯品会网站上抓取价格标签。",
            requirement_type="sop",
        )
        profile["stage"] = "confirm"
        result = apply_answer(profile, "确认需求画像")

        self.assertEqual(result["reason"], "confirmation_blocked")
        self.assertEqual(result["profile"]["stage"], "clarifying")
        self.assertNotEqual(result["profile"]["stage"], "done")

    def test_ai_facts_write_to_the_matching_business_context_fields(self) -> None:
        profile = new_profile(
            title="售后退款自动审核",
            department="客服中心",
            summary="根据订单、物流和退款原因自动审核退款。",
            requirement_type="sop",
        )
        profile = apply_ai_direction(
            profile,
            {
                "status": "continue",
                "requirementType": "sop",
                "questionKey": "successMetric",
                "question": "上线后用什么指标衡量效果？",
                "facts": {
                    "systemLandscape": "订单系统、物流系统和客服工单系统",
                    "dataDefinition": "退款金额以实际支付金额为准",
                    "approvalBoundary": "高金额和疑似欺诈转人工",
                },
                "criticalGaps": [],
            },
        )
        self.assertEqual(
            profile["businessContext"]["systemLandscape"],
            "订单系统、物流系统和客服工单系统",
        )
        self.assertEqual(
            profile["businessContext"]["dataDefinition"],
            "退款金额以实际支付金额为准",
        )
        self.assertEqual(
            profile["businessContext"]["approvalBoundary"],
            "高金额和疑似欺诈转人工",
        )

    def test_ai_single_question_truncates_compound_prompt(self) -> None:
        self.assertEqual(
            _single_question("哪些情况转人工？另外，成功指标是什么？"),
            "哪些情况转人工？",
        )
        self.assertEqual(
            _single_question("哪些情况转人工，以及由谁审批"),
            "哪些情况转人工？",
        )

    def test_streaming_json_question_can_be_read_before_json_is_complete(self) -> None:
        partial = '{"status":"continue","question":"请说明当前由谁在商品后台'
        self.assertEqual(
            _partial_json_string_field(partial, "question"),
            "请说明当前由谁在商品后台",
        )

    def test_uploaded_materials_are_distinguished_from_unreadable_images(self) -> None:
        text = _project_materials_text(
            [
                {
                    "original_filename": "库存口径.xlsx",
                    "parse_status": "parsed",
                    "extracted_text": "[工作表：库存]\nSKU | 可售库存\nA-001 | 12",
                },
                {
                    "original_filename": "现场截图.png",
                    "parse_status": "stored",
                    "extracted_text": "",
                },
            ]
        )
        self.assertIn("A-001 | 12", text)
        self.assertIn("现场截图.png", text)
        self.assertIn("不能据此推断具体事实", text)

    def test_ai_judged_profile_is_high_confidence_without_checklist_gaps(self) -> None:
        profile = new_profile(
            title="售后退款自动审核",
            department="客服中心",
            summary="根据订单、物流和退款原因自动审核退款。",
            requirement_type="sop",
        )
        profile["stage"] = "done"
        profile["aiJudged"] = True
        profile["aiGaps"] = []
        assessment = assess(profile, [])
        completeness = next(item for item in assessment["dimensions"] if item["name"] == "信息完整度")
        self.assertEqual(assessment["confidence"], "高")
        self.assertEqual(completeness["score"], 100)

    def test_model_assessment_can_override_rigid_rule_priority_with_evidence(self) -> None:
        profile = new_profile(
            title="库存补货与断货预警",
            department="供应链",
            summary="影响多渠道补货和库存连续性。",
            requirement_type="sop",
        )
        fallback = assess(profile, [])
        merged = merge_model_assessment(
            fallback,
            {
                "model": "grok-test",
                "value_score": 96,
                "priority": "P0",
                "business_criticality": "核心经营主流程",
                "risk_level": "高",
                "confidence": "高",
                "summary": "影响核心供应链连续经营，存在断货与积压风险。",
                "rationale": ["覆盖补货主流程", "断货和积压会直接影响销售与资金占用"],
                "evidence_gaps": [],
                "dimensions": [
                    {"name": "供应连续性", "score": 95, "evidence": "影响核心渠道补货"},
                ],
            },
        )
        self.assertEqual(merged["priority"], "P0")
        self.assertEqual(merged["value_score"], 96)
        self.assertEqual(merged["assessment_source"], "llm")
        self.assertEqual(merged["business_criticality"], "核心经营主流程")

    def test_model_p0_with_low_confidence_is_downgraded(self) -> None:
        fallback = assess(
            new_profile(
                title="库存需求",
                department="供应链",
                summary="库存需求",
                requirement_type="sop",
            ),
            [],
        )
        merged = merge_model_assessment(
            fallback,
            {
                "value_score": 95,
                "priority": "P0",
                "risk_level": "高",
                "confidence": "低",
                "dimensions": [],
            },
        )
        self.assertEqual(merged["priority"], "P1")

    def test_model_assessment_json_is_normalized(self) -> None:
        result = _parse_requirement_assessment(
            '{"valueScore": 88, "priority": "P1", '
            '"businessCriticality": "核心经营主流程", "riskLevel": "中", '
            '"confidence": "高", "summary": "影响补货连续性", '
            '"rationale": ["覆盖核心渠道"], "evidenceGaps": ["缺少库存金额"], '
            '"dimensions": [{"name": "库存与资金风险", "score": 82, "evidence": "存在积压风险"}]}',
            "test-model",
        )
        self.assertEqual(result["value_score"], 88)
        self.assertEqual(result["priority"], "P1")
        self.assertEqual(result["model"], "test-model")
        self.assertEqual(result["evidence_gaps"], ["缺少库存金额"])


class PlanTests(unittest.TestCase):
    def test_ai_research_workbook_uses_business_sheet_and_pipe_delimited_cells(self) -> None:
        text = """[工作表：SOP流程拆解]
SOP名称 |  | 采购订单（返单）
流程目的 |  | 按计划交期完成大货入仓，避免断货
SOP完整文字描述 |  | 采购在金蝶云推单，经理审核后通知供应商
序号 | 具体动作 | 输入（需要什么） | 步骤类型
1 | 根据采购申请单在金蝶云推单 | 采购申请表 | 数据信息处理
2 | 采购经理审核采购订单 | 采购订单 | 分析决策

[工作表：类型判断指南]
SOP名称 |  | 示例流程，不应进入业务画像
序号 | 具体动作
1 | 登录生意参谋并下载示例日报
"""
        profile = add_material_candidates(
            new_profile(title="品牌采购", department="创研", summary="", requirement_type="sop"),
            source_id="research-1",
            filename="品牌采购.xlsx",
            parser_type="xlsx",
            text=text,
        )
        candidates = {item["field"]: item for item in profile["materialCandidates"]}
        self.assertEqual(candidates["businessObject"]["value"], "采购订单（返单）")
        self.assertIn("避免断货", candidates["purpose"]["value"])
        self.assertIn("金蝶云推单", candidates["currentProcess"]["value"])
        self.assertEqual(len(candidates["steps"]["value"]), 2)
        self.assertNotIn("示例日报", str(candidates))
        self.assertEqual(candidates["steps"]["locator"], "工作表：SOP流程拆解")

    def test_material_candidate_does_not_become_confirmed_before_acceptance(self) -> None:
        profile = new_profile(
            title="库存设置", department="运营", summary="", requirement_type="sop"
        )
        proposed = add_material_candidates(
            profile,
            source_id="file-1",
            filename="库存教程.xlsx",
            parser_type="xlsx",
            text="[工作表：配置]\n业务对象：直营网店 SKU 库存\n人工边界：负库存转人工",
        )
        self.assertEqual(proposed["businessContext"]["businessObject"], "")
        self.assertFalse(proposed.get("sourceEvidence"))
        self.assertTrue(any(item["field"] == "businessObject" for item in proposed["materialCandidates"]))
        self.assertEqual(confirmation_gaps_for(proposed)[0]["key"], "materialReview")

    def test_material_candidates_detect_conflict_and_accept_one_source(self) -> None:
        profile = new_profile(
            title="库存设置", department="运营", summary="", requirement_type="sop"
        )
        profile = add_material_candidates(
            profile,
            source_id="file-1",
            filename="旧流程.xlsx",
            parser_type="xlsx",
            text="[工作表：配置]\n数据口径：可售库存等于实物库存减锁定库存",
        )
        profile = add_material_candidates(
            profile,
            source_id="file-2",
            filename="新流程.docx",
            parser_type="docx",
            text="[段落 8]\n数据口径：可售库存等于实物库存减锁定库存减预售占用",
        )
        candidates = [item for item in profile["materialCandidates"] if item["field"] == "dataDefinition"]
        self.assertEqual(len(candidates), 2)
        self.assertTrue(all(item["status"] == "conflict" for item in candidates))

        resolved = resolve_material_candidate(
            profile, candidate_id=candidates[1]["id"], action="accept"
        )
        self.assertIn("预售占用", resolved["businessContext"]["dataDefinition"])
        self.assertEqual(candidates[0]["status"], "conflict")
        statuses = {
            item["source"]: item["status"]
            for item in resolved["materialCandidates"]
            if item["field"] == "dataDefinition"
        }
        self.assertEqual(statuses["新流程.docx"], "accepted")
        self.assertEqual(statuses["旧流程.xlsx"], "superseded")
        self.assertEqual(resolved["sourceEvidence"][-1]["status"], "accepted")

    def test_material_candidate_can_be_edited_or_rejected(self) -> None:
        profile = add_material_candidates(
            new_profile(title="库存设置", department="运营", summary="", requirement_type="sop"),
            source_id="file-1",
            filename="教程.xls",
            parser_type="xls",
            text="[工作表：流程]\n适用范围：全部商品\n成功指标：处理完成",
        )
        scope = next(item for item in profile["materialCandidates"] if item["field"] == "scope")
        metric = next(item for item in profile["materialCandidates"] if item["field"] == "successMetric")
        profile = resolve_material_candidate(
            profile,
            candidate_id=scope["id"],
            action="edit",
            edited_value="直营网店常规在售商品",
        )
        profile = resolve_material_candidate(profile, candidate_id=metric["id"], action="reject")
        self.assertEqual(profile["businessContext"]["scope"], "直营网店常规在售商品")
        self.assertEqual(profile["businessContext"]["successMetric"], "")
        self.assertEqual(
            next(item for item in profile["materialCandidates"] if item["id"] == metric["id"])["status"],
            "rejected",
        )

    def test_generated_plan_can_be_exported_as_docx(self) -> None:
        payload = {
            "title": "库存半自动产品方案",
            "version": 2,
            "markdown_content": "# 方案目标\n\n- **系统负责**：生成库存建议\n- 人工确认后执行",
        }
        stream = render_document_docx(payload, {"title": "库存设置优化"})
        self.assertTrue(stream.getvalue().startswith(b"PK"))
        parsed = Document(BytesIO(stream.getvalue()))
        text = "\n".join(item.text for item in parsed.paragraphs)
        self.assertIn("库存半自动产品方案", text)
        self.assertIn("系统负责", text)

    def test_uploaded_material_builds_profile_with_source_locator(self) -> None:
        profile = new_profile(
            title="库存设置优化",
            department="商品运营",
            summary="",
            requirement_type="sop",
        )
        enriched = merge_material_into_profile(
            profile,
            filename="后台库存设置教程.xls",
            parser_type="xls",
            text=(
                "[工作表：操作教程]\n"
                "业务对象：直营网店商品库存\n"
                "适用范围：常规在售 SKU，不含预售商品\n"
                "当前处理方式：运营从 ERP 导出库存后登录商品后台逐条设置\n"
                "涉及系统：金蝶云、百胜 E3、电商后台\n"
                "人工边界：缺少仓库映射或库存为负时转人工\n"
                "成功指标：库存设置准确率达到 99%\n"
                "流程目的：减少重复录入并避免超卖\n"
                "1. 从金蝶云导出可用库存\n"
                "2. 按仓库和 SKU 映射转换\n"
                "3. 在商品后台提交库存"
            ),
        )
        self.assertEqual(enriched["businessContext"]["businessObject"], "直营网店商品库存")
        self.assertEqual(len(enriched["sop"]["steps"]), 3)
        evidence = enriched["sourceEvidence"]
        self.assertTrue(any(item["source"] == "后台库存设置教程.xls" for item in evidence))
        self.assertTrue(any(item["locator"] == "工作表：操作教程" for item in evidence))

    def test_ragflow_failure_keeps_local_knowledge_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            Path(directory, "inventory.md").write_text(
                "库存设置需要校验 SKU 与仓库映射。", encoding="utf-8"
            )
            settings = Settings(
                knowledge_local_path=directory,
                ragflow_enabled=True,
                ragflow_api_key="test-key",
                ragflow_base_url="http://ragflow.invalid",
            )
            with patch("app.knowledge.httpx.post", side_effect=httpx.ConnectError("offline")):
                references = collect_plan_references(
                    settings=settings,
                    project={"title": "库存设置", "summary": "SKU 仓库映射"},
                    profile=new_profile(
                        title="库存设置", department="运营", summary="SKU 仓库映射", requirement_type="sop"
                    ),
                    source_files=[],
                )
            self.assertTrue(any(item["provider"] == "local" for item in references))

    def test_ragflow_chunks_are_normalized_and_combined_with_local_retrieval(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            Path(directory, "inventory.md").write_text("库存需要校验 SKU 映射。", encoding="utf-8")
            settings = Settings(
                knowledge_local_path=directory,
                ragflow_enabled=True,
                ragflow_api_key="test-key",
                ragflow_base_url="http://ragflow.test",
            )
            response = Mock()
            response.raise_for_status.return_value = None
            response.json.return_value = {
                "data": {"chunks": [{"content": "E3 库存接口支持库存查询", "docnm_kwd": "E3接口.pdf", "similarity": 0.91}]}
            }
            with patch("app.knowledge.httpx.post", return_value=response):
                references = collect_plan_references(
                    settings=settings,
                    project={"title": "库存设置", "summary": "SKU 映射"},
                    profile=new_profile(title="库存设置", department="运营", summary="SKU 映射", requirement_type="sop"),
                    source_files=[],
                )
            providers = {item["provider"] for item in references}
            self.assertIn("ragflow", providers)
            self.assertIn("local", providers)

    def test_inventory_case_reaches_four_solution_documents(self) -> None:
        profile = new_profile(
            title="后台库存自动设置",
            department="商品运营",
            summary="根据 ERP 和 E3 库存更新商品后台。",
            requirement_type="sop",
        )
        profile = merge_material_into_profile(
            profile,
            filename="后台库存设置教程.xls",
            parser_type="xls",
            text=(
                "[工作表：库存流程]\n业务对象：商品 SKU 可售库存\n"
                "适用范围：直营网店常规商品\n"
                "当前处理方式：运营每天导出 ERP 库存并录入电商后台\n"
                "数据口径：可售库存等于实物库存减锁定库存\n"
                "涉及系统：金蝶云、百胜 E3、电商后台\n"
                "人工边界：负库存、映射缺失和接口失败转人工\n"
                "成功指标：准确率 99%，处理时间低于 30 分钟\n"
                "流程目的：减少超卖与重复录入\n"
                "1. 导出库存\n2. 校验 SKU 和仓库映射\n3. 更新商品后台\n4. 输出异常清单"
            ),
        )
        profile["sop"]["owner"] = "商品运营主管"
        profile["stage"] = "done"
        documents = generate_documents(
            {"title": "后台库存自动设置", "summary": "库存同步"}, profile, []
        )
        self.assertEqual(len(documents), 4)
        self.assertEqual({item["tier"] for item in documents}, {"semi_auto", "full_auto"})
        self.assertTrue(all("库存" in item["markdown_content"] for item in documents))
    def test_excel_and_image_materials_are_supported(self) -> None:
        self.assertEqual(parser_type_for("业务事实.xlsx"), "xlsx")
        self.assertEqual(parser_type_for("现场截图.png"), "png")

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "业务事实.xlsx"
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "库存口径"
            sheet.append(["SKU", "可售库存", "处理方式"])
            sheet.append(["A-001", 12, "继续销售"])
            workbook.save(workbook_path)

            parsed_workbook = parse_document(workbook_path, workbook_path.name)
            self.assertEqual(parsed_workbook["parser_type"], "xlsx")
            self.assertIn("库存口径", parsed_workbook["text"])
            self.assertIn("A-001 | 12 | 继续销售", parsed_workbook["text"])

            image_path = root / "现场截图.png"
            image_path.write_bytes(
                base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
                    "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
                )
            )
            parsed_image = parse_document(image_path, image_path.name)
            self.assertEqual(parsed_image["parser_type"], "png")
            self.assertEqual(parsed_image["text"], "")
            self.assertEqual(parsed_image["metadata"]["textExtraction"], "pending_vision_ocr")

    def test_explanation_precedes_four_solution_documents(self) -> None:
        project = {
            "id": "project-1",
            "title": "自动设置库存",
        }
        profile = new_profile(
            title="自动设置库存",
            department="商品运营",
            summary="根据库存状态生成后台库存设置建议。",
            requirement_type="decision",
        )
        profile["stage"] = "done"
        profile["businessContext"].update(
            {
                "businessChain": "inventory_operations",
                "businessObject": "指定 SKU 的可售库存设置",
                "scope": "先覆盖主店常规商品",
                "currentProcess": "运营每天导出库存表后手工设置商品后台库存",
                "dataDefinition": "可售库存等于实物库存减锁定库存",
                "systemLandscape": "商品后台和库存报表",
                "approvalBoundary": "库存异常时由运营主管确认",
                "successMetric": "库存设置时效缩短且错误率下降",
            }
        )
        profile["decision"]["purpose"] = "减少人工设置库存的时间和错误"

        explanation = generate_requirement_explanation(project, profile, [])
        solutions = generate_documents(project, profile, [])

        self.assertEqual(explanation["tier"], "requirement_explanation")
        self.assertEqual(explanation["document_type"], "explanation")
        self.assertIn("不是最终需求方案", explanation["markdown_content"])
        self.assertEqual(explanation["prompt_version"], "plans-v8")
        self.assertEqual(len(solutions), 4)
        self.assertEqual(
            {(item["tier"], item["document_type"]) for item in solutions},
            {
                ("semi_auto", "product"),
                ("semi_auto", "tech"),
                ("full_auto", "product"),
                ("full_auto", "tech"),
            },
        )
        self.assertNotIn("价值评分", "\n".join(item["markdown_content"] for item in solutions))
        self.assertTrue(all("## 选型结论" in item["markdown_content"] for item in solutions))
        self.assertTrue(all("## 方案摘要" not in item["markdown_content"] for item in solutions))

    def test_technical_plans_follow_demo_execution_outline(self) -> None:
        project = {
            "id": "project-rpa",
            "title": "库存后台自动设置",
        }
        profile = new_profile(
            title=project["title"],
            department="商品运营",
            summary="按库存结果生成后台设置建议。",
            requirement_type="decision",
        )
        profile["stage"] = "done"
        profile["businessContext"].update(
            {
                "systemLandscape": "WMS仓储、商品后台和 Excel",
                "dataDefinition": "可售库存 = 实物库存 - 锁定库存",
                "approvalBoundary": "异常库存由运营主管确认",
            }
        )
        profile["decision"].update(
            {
                "inputs": [
                    {"name": "库存数据", "source": "WMS仓储", "method": "RPA抓取"},
                    {"name": "人工调整表", "source": "Excel", "method": "定时导出"},
                ],
                "rules": "按库存阈值生成设置建议",
                "actions": [
                    {
                        "result": "设置库存",
                        "who": "运营",
                        "what": "确认后在商品后台执行",
                        "where": "商品后台",
                    }
                ],
            }
        )

        documents = generate_documents(project, profile, [])
        technical = [item for item in documents if item["document_type"] == "tech"]

        self.assertEqual(len(technical), 2)
        for document in technical:
            content = document["markdown_content"]
            for heading in (
                "## 数据流",
                "## 集成点",
                "## 依赖与前置",
                "## 环境变量",
                "## 测试计划",
                "## 完成定义 DoD",
            ):
                self.assertIn(heading, content)
            self.assertIn("RPA", content)
            self.assertIn("定时任务", content)
            self.assertIn("商品后台", content)

    def test_plans_use_business_facts_instead_of_generic_template_copy(self) -> None:
        project = {
            "id": "project-inventory",
            "title": "库存周转控制",
            "summary": "每周根据销售、库存、在途和采购周期生成备货需求，降低库存风险。",
        }
        profile = new_profile(
            title=project["title"],
            department="计划部",
            summary=project["summary"],
            requirement_type="sop",
        )
        profile["stage"] = "done"
        profile["businessContext"].update(
            {
                "systemLandscape": "百胜E3、各电商后台、线下Excel",
            }
        )
        references = [
            {
                "sourceType": "project_material",
                "source": "周预测表.xlsx",
                "excerpt": (
                    "平台库存/平台在途、E3库存/移仓未入库、近7天/近30天销量、销售计划、"
                    "采购在途。库存+在途+预计回货计划-安全库存-月销售量；"
                    "（库存+在途）/月销=周转(月)。输出备货计划和采购申请。"
                ),
            }
        ]

        documents = generate_documents(project, profile, [], references)
        product = next(
            item
            for item in documents
            if item["tier"] == "semi_auto" and item["document_type"] == "product"
        )["markdown_content"]
        tech = next(
            item
            for item in documents
            if item["tier"] == "semi_auto" and item["document_type"] == "tech"
        )["markdown_content"]

        self.assertIn("库存周转", product)
        self.assertIn("备货数量建议", product)
        self.assertIn("库存 + 在途", product)
        self.assertIn("各系统数据 -> 脚本拉取/读取 -> 清洗合并 -> 计算指标", tech)
        self.assertIn("企业微信/钉钉机器人", tech)
        self.assertNotIn("系统负责自动触发、处理并回写", product)
        self.assertNotIn("具体页面、接口、权限和上线范围由评审人补齐", tech)
        full_auto_product = next(
            item
            for item in documents
            if item["tier"] == "full_auto" and item["document_type"] == "product"
        )["markdown_content"]
        self.assertIn("规则未覆盖", full_auto_product)
        self.assertNotIn("人工保留**：计划部：下采购申请单", full_auto_product)

    def test_plan_references_keep_materials_and_knowledge_separate(self) -> None:
        project = {
            "id": "project-2",
            "title": "自动设置库存",
            "summary": "商品运营希望根据 WMS 库存和商品后台规则生成库存设置建议。",
        }
        profile = new_profile(
            title="自动设置库存",
            department="商品运营",
            summary=project["summary"],
            requirement_type="sop",
        )
        profile["businessContext"].update(
            {
                "businessChain": "inventory_operations",
                "businessObject": "SKU 可售库存",
            }
        )
        references = [
            {
                "sourceType": "project_material",
                "provider": "uploaded_material",
                "source": "库存口径.md",
                "excerpt": "可售库存不包含锁定库存和残次品。",
                "confidence": "fact",
            },
            {
                "sourceType": "knowledge_base",
                "provider": "local",
                "source": "qibang-reference-scenarios.md",
                "excerpt": "WMS、商品后台和飞书是历史试点中出现过的参考系统。",
                "confidence": "reference",
            },
        ]

        explanation = generate_requirement_explanation(project, profile, [], references)
        solutions = generate_documents(project, profile, [], references)

        self.assertEqual(explanation["knowledge_refs"], references)
        self.assertNotIn("当前项目材料（事实依据）", explanation["markdown_content"])
        self.assertNotIn("七邦参考知识（仅供参考，不等同于当前事实）", explanation["markdown_content"])
        self.assertTrue(all(item["knowledge_refs"] == references for item in solutions))
        self.assertTrue(all("AI 推测与落地边界" not in item["markdown_content"] for item in solutions))

        requester_view = _requester_document(
            {
                **solutions[0],
                "model_name": "hidden-model",
                "knowledge_refs": references,
            }
        )
        self.assertNotIn("qibang-reference-scenarios.md", requester_view["markdown_content"])
        self.assertNotIn("历史试点中出现过", requester_view["markdown_content"])
        self.assertNotIn("内部参考边界", requester_view["markdown_content"])
        self.assertNotIn("knowledge_refs", requester_view)

    def test_local_knowledge_reference_is_not_treated_as_project_fact(self) -> None:
        settings = Settings(
            knowledge_local_path="knowledge",
            ragflow_enabled=False,
        )
        project = {
            "title": "自动上下架",
            "summary": "商品运营希望结合 SKU、库存和商品后台状态处理上下架。",
        }
        profile = new_profile(
            title=project["title"],
            department="商品运营",
            summary=project["summary"],
            requirement_type="sop",
        )
        profile["businessContext"]["businessChain"] = "inventory_operations"

        references = collect_plan_references(
            settings=settings,
            project=project,
            profile=profile,
            source_files=[],
        )

        self.assertTrue(any(item["sourceType"] == "knowledge_base" for item in references))
        self.assertTrue(all(item["confidence"] == "reference" for item in references))

    def test_reviewer_uploaded_knowledge_is_used_as_reference(self) -> None:
        settings = Settings(
            knowledge_local_path="missing-knowledge",
            ragflow_enabled=False,
        )
        project = {
            "title": "自动上下架",
            "summary": "商品运营希望结合 SKU、库存和商品后台状态处理上下架。",
        }
        profile = new_profile(
            title=project["title"],
            department="商品运营",
            summary=project["summary"],
            requirement_type="sop",
        )
        profile["businessContext"]["businessChain"] = "inventory_operations"

        references = collect_plan_references(
            settings=settings,
            project=project,
            profile=profile,
            source_files=[],
            knowledge_files=[
                {
                    "id": "knowledge-1",
                    "original_filename": "商品后台操作规范.pdf",
                    "extracted_text": "商品后台上下架需要校验 SKU 状态、库存和审批边界。",
                }
            ],
        )

        uploaded = [
            item
            for item in references
            if item.get("provider") == "reviewer_upload"
        ]
        self.assertEqual(len(uploaded), 1)
        self.assertEqual(uploaded[0]["sourceType"], "knowledge_base")
        self.assertEqual(uploaded[0]["confidence"], "reference")


if __name__ == "__main__":
    unittest.main()
